"""Outputs people can take away: Excel workbooks that explain themselves, calendar files, templates.

Every workbook starts with a "Read me" sheet: what it is, when it was produced, live or demo, where the numbers
come from, and what each column means. Header rows are bold and frozen, columns sized to their content.
Callers pass only rows the person may see (scope is checked in services before calling these).
"""
from __future__ import annotations

import io
from datetime import datetime, timedelta

import numpy as np
import pandas as pd

from .config import day_label, queue_names, shift_window
from .schemas import DOMAINS


def workbook(title: str, sheets: dict[str, pd.DataFrame], *, context: str, source: str,
             definitions: dict[str, str] | None = None, kind: str = "planned") -> bytes:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    readme = [("What this is", title), ("Produced", datetime.now().strftime("%d %b %Y %H:%M")),
              ("Workspace", context), ("Figures are", kind), ("Source", source),
              ("Note", "All people, clients and figures in this proof of concept are simulated.")]
    for k, v in (definitions or {}).items():
        readme.append((k, v))
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        pd.DataFrame(readme, columns=["Item", "Explanation"]).to_excel(xw, sheet_name="Read me", index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="E6EFF6")
            for i, col in enumerate(ws.columns, start=1):
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    return out.getvalue()


def _ctx(w) -> str:
    return "Demo scenario (simulated)" if getattr(w, "scenario", False) else "Live workspace"


SHIFT_DEF = {"Shift": "Morning 06–15, Day 09–18, Evening 14–23, Night 22–07 (9 h incl. 1 h break, 8 paid)",
             "Queue": "The type of call worked (process × language)",
             "Extension": "Hours added before or after the shift (overtime)"}


def roster_rows(w, agent_ids: set | None = None) -> pd.DataFrame:
    names = w.data["agents"].set_index("agent_id")["name"].to_dict()
    teams = w.data["agents"].set_index("agent_id")["team"].to_dict()
    ro = w.roster if agent_ids is None else w.roster[w.roster.agent_id.isin(agent_ids)]
    rows = []
    for r in ro.sort_values(["day", "agent_id"]).itertuples():
        s, e = shift_window(w.policy, r.day, r.shift, r.ext_before, r.ext_after)
        rows.append({"Date": day_label(w.policy, r.day), "Agent": names.get(r.agent_id, r.agent_id),
                     "Employee ID": r.agent_id, "Team": teams.get(r.agent_id, ""), "Shift": r.shift,
                     "Starts": f"{int(s % 24):02d}:{int(round((s % 1) * 60)):02d}",
                     "Ends": f"{int(e % 24):02d}:{int(round((e % 1) * 60)):02d}", "Queue": r.queue,
                     "Extension (h)": r.ext_before + r.ext_after, "Change type": r.kind if r.kind != "plan" else ""})
    return pd.DataFrame(rows)


def roster_grid(w, agent_ids: set | None = None) -> pd.DataFrame:
    """Agents × days, cells like 'M · Claims-EN' (blank = day off)."""
    ag = w.data["agents"] if agent_ids is None else w.data["agents"][w.data["agents"].agent_id.isin(agent_ids)]
    days = range(w.policy["horizon"]["days"])
    cell = {(r.agent_id, r.day): f"{r.shift[0]} · {r.queue}" for r in w.roster.itertuples()}
    leave = {(r.agent_id, r.day): "Leave" for r in w.data["leave"].itertuples()}
    rows = []
    for a in ag.sort_values(["team", "agent_id"]).itertuples():
        row = {"Employee ID": a.agent_id, "Agent": a.name, "Team": a.team}
        for d in days:
            row[day_label(w.policy, d)] = cell.get((a.agent_id, d)) or leave.get((a.agent_id, d), "")
        rows.append(row)
    return pd.DataFrame(rows)


def coverage_rows(w, days: list[int] | None = None) -> pd.DataFrame:
    from .roster import coverage
    cov = coverage(w.policy, w.roster)
    per_day = w.req.shape[1] // w.policy["horizon"]["days"]
    rows = []
    for d in days if days is not None else range(w.policy["horizon"]["days"]):
        for qi, q in enumerate(queue_names(w.policy)):
            sl = slice(d * per_day, (d + 1) * per_day)
            need, have = w.req[qi, sl], cov[qi, sl]
            short = np.maximum(need - have, 0)
            calls = float(w.fc["mid"][qi, sl].sum()) if w.fc else None
            rows.append({"Date": day_label(w.policy, d), "Queue": q, "Calls forecast": round(calls) if calls else None,
                         "Required (people-h)": need.sum() / 2, "Scheduled (people-h)": np.minimum(have, need).sum() / 2,
                         "Short (people-h)": short.sum() / 2,
                         "Covered %": round(100 * (1 - short.sum() / need.sum()), 1) if need.sum() else 100.0})
    return pd.DataFrame(rows)


def roster_workbook(w, agent_ids: set | None = None, title: str = "Published roster") -> bytes:
    return workbook(title, {"By agent and day": roster_grid(w, agent_ids), "Shift list": roster_rows(w, agent_ids)},
                    context=_ctx(w), source="Live roster (published plan plus approved changes)", kind="planned",
                    definitions=SHIFT_DEF)


def team_pack(w, team: str) -> bytes:
    from .report import weekly_report
    ids = set(w.data["agents"][w.data["agents"].team == team].agent_id)
    rep = weekly_report(w, w.today // 7)
    att = rep["attendance"][rep["attendance"].Team == team]
    return workbook(f"Team {team} pack", {"Schedule": roster_grid(w, ids), "Shift list": roster_rows(w, ids),
                                          "Attendance this week": att,
                                          "Requests": pd.DataFrame([r for r in rep["requests"].to_dict("records")
                                                                    if r["Agent"] in ids])},
                    context=_ctx(w), source="Live roster and recorded attendance", kind="planned and recorded",
                    definitions={**SHIFT_DEF, "Attendance": "Recorded events only; no record is never assumed worked"})


def ops_brief(w) -> bytes:
    from . import priority
    days = list(range(w.today, min(w.policy["horizon"]["days"], w.today + 3)))
    cov = coverage_rows(w, days)
    open_inc = [i for i in w.incidents if i.status not in ("closed", "no_gap", "fixed")]
    risks = []
    for i in open_inc:
        s = priority.assess(w, i)
        risks.append({"Alert": i.id, "Date": day_label(w.policy, i.day), "Team": i.owner_team, "Issue": i.title,
                      "Severity": s["severity"], "Gap (people-h)": s["short_ph"], "Status": i.status})
    return workbook("Daily operations brief (next 72 h)", {
        "Coverage next 3 days": cov, "Open alerts": pd.DataFrame(risks),
        "Watch list (<95%)": cov[cov["Covered %"] < 95]},
        context=_ctx(w), source="Live roster against forecast requirement (Erlang C, configured shrinkage)",
        kind="planned against forecast",
        definitions={"Covered %": "Required people-time the roster provides; forecast-based, not measured service",
                     "Gap (people-h)": "One person short for one hour = 1"})


def forecast_workbook(w) -> bytes:
    per_day = w.req.shape[1] // w.policy["horizon"]["days"]
    rows = []
    from .roster import coverage
    cov = coverage(w.policy, w.roster)
    for qi, q in enumerate(queue_names(w.policy)):
        for t in range(w.req.shape[1]):
            if not w.req[qi, t] and not (w.fc and w.fc["mid"][qi, t]):
                continue
            rows.append({"Date": day_label(w.policy, t // per_day), "Time": f"{(t % per_day) // 2:02d}:{30 * (t % 2):02d}",
                         "Queue": q, "Calls forecast": round(float(w.fc["mid"][qi, t]), 1) if w.fc else None,
                         "Calls low": round(float(w.fc["low"][qi, t]), 1) if w.fc else None,
                         "Calls high": round(float(w.fc["high"][qi, t]), 1) if w.fc else None,
                         "People required": int(w.req[qi, t]), "People scheduled": int(cov[qi, t])})
    return workbook("Forecast and staffing requirement per half hour", {"Per half hour": pd.DataFrame(rows),
                                                                       "Per day and queue": coverage_rows(w)},
                    context=_ctx(w), source="Seasonal forecast from call history; Erlang C; shrinkage from rules",
                    kind="forecast (prediction, not actual)",
                    definitions={"Calls low / high": "About an 80% band around the forecast"})


def reference_template(domain: str, records: list[dict]) -> bytes:
    """Current reference data as a ready-filled file with friendly headers, to edit and upload back."""
    d = DOMAINS[domain]
    cols = [f for f in d.fields if not (domain == "volumes" and f.name == "time")]
    rows = []
    for r in records:
        row = {}
        for f in cols:
            v = r.get(f.name)
            if isinstance(v, list):
                v = (f"{round(v[0] * 100)}% in {v[1]:g}s" if f.kind == "sla" else
                     f"{int(v[0]):02d}:{int(round(v[0] % 1 * 60)):02d}-{int(v[1]):02d}:{int(round(v[1] % 1 * 60)):02d}"
                     if f.kind == "hours" else "; ".join(map(str, v)))
            if f.kind == "weekday" and isinstance(v, int):
                v = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"][v]
            if f.kind == "bool" and v is not None:
                v = "Yes" if v else "No"
            row[f.label + (" *" if f.required else "")] = v
        rows.append(row)
    df = pd.DataFrame(rows, columns=[f.label + (" *" if f.required else "") for f in cols])
    guide = pd.DataFrame([{"Column": f.label, "Mandatory": "Yes" if f.required else "No", "Help": f.help or
                           {"date": "Any common date format; day first (12/10/2026 = 12 October)",
                            "queues": "Separate several with ; or ,", "bool": "Yes / No", "sla": "e.g. 80/20 or 80% in 20s",
                            "hours": "e.g. 09:00-21:00 or 24x7", "duration_sec": "e.g. 6 min or 360 s",
                            "weekday": "e.g. Saturday"}.get(f.kind, "")} for f in cols])
    return workbook(f"{d.label}: current data (edit and upload back)", {d.label[:31]: df, "Columns": guide},
                    context="Reference data", source="ShiftSense reference data, latest versions", kind="reference data",
                    definitions={"How to use": "Edit or add rows, keep the headers, upload in Data hub → Add or update. "
                                               "Only changed rows are applied. Columns marked * are mandatory."})


def change_log(rows: list[dict]) -> bytes:
    df = pd.DataFrame([{"Change": f"#{r['id']}", "Data": DOMAINS[r["domain"]].label, "Status": r["status"],
                        "Submitted by": r["submitted_by"], "Submitted": r["submitted_at"],
                        "Decided by": r["decided_by"] or "", "Effective from": r["effective_from"],
                        "Reason": r["reason"] or "", "Source": r["source"], "Converted by": r["converted_by"],
                        "New": (r["summary"] or {}).get("new", 0), "Changed": (r["summary"] or {}).get("changed", 0),
                        "Removed": (r["summary"] or {}).get("delete", 0)} for r in rows])
    return workbook("Reference data change log", {"Changes": df}, context="Reference data",
                    source="Change sets (append-only)", kind="recorded")


def ics(w, agent_id: str) -> bytes:
    """The agent's shifts as a calendar file for phone or Outlook (times in the operation's local time)."""
    from .refdata import start_date
    start = start_date(w.policy)
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//ShiftSense//PoC//EN", "CALSCALE:GREGORIAN"]
    for r in w.roster[w.roster.agent_id == agent_id].sort_values("day").itertuples():
        s, e = shift_window(w.policy, r.day, r.shift, r.ext_before, r.ext_after)
        base = datetime.combine(start, datetime.min.time())
        ds, de = base + timedelta(hours=float(s)), base + timedelta(hours=float(e))
        lines += ["BEGIN:VEVENT", f"UID:{w.run_id}-{agent_id}-{r.day}@shiftsense",
                  f"DTSTAMP:{datetime.now():%Y%m%dT%H%M%S}", f"DTSTART:{ds:%Y%m%dT%H%M%S}", f"DTEND:{de:%Y%m%dT%H%M%S}",
                  f"SUMMARY:{r.shift} shift · {r.queue}", "DESCRIPTION:From ShiftSense (simulated PoC data)",
                  "END:VEVENT"]
    lines.append("END:VCALENDAR")
    return "\r\n".join(lines).encode()


def client_service_report(w) -> tuple[pd.DataFrame, bytes]:
    """Per client contract and day: planned coverage against the requirement, and, where imported, actual calls,
    forecast error and service level achieved."""
    from . import refdata
    cov = coverage_rows(w)
    client = {q: (w.policy["queues"][q].get("client") or "Unassigned") for q in queue_names(w.policy)}
    cov["Client"] = cov["Queue"].map(client)
    sla = {q: f"{w.policy['queues'][q]['sla_pct'] * 100:g}% in {w.policy['queues'][q]['sla_sec']:g}s"
           for q in queue_names(w.policy)}
    cov["Service level promised"] = cov["Queue"].map(sla)
    act = refdata.actuals(w.policy) if getattr(w, "ref_linked", False) else pd.DataFrame()
    fa = refdata.forecast_accuracy(w.policy, w.fc, act) if len(act) else pd.DataFrame()
    sv = refdata.actual_service(act) if len(act) else pd.DataFrame()
    if len(fa):
        fa = fa.assign(Date=[day_label(w.policy, d) for d in fa.day]).rename(
            columns={"queue": "Queue", "actual": "Actual calls", "error_pct": "Forecast error %"})
        cov = cov.merge(fa[["Date", "Queue", "Actual calls", "Forecast error %"]], on=["Date", "Queue"], how="left")
    if len(sv):
        sv = sv.assign(Date=[day_label(w.policy, d) for d in sv.day]).rename(
            columns={"queue": "Queue", "service_level_pct": "Service level achieved %"})
        cov = cov.merge(sv[["Date", "Queue", "Service level achieved %"]], on=["Date", "Queue"], how="left")
    cols = ["Client", "Queue", "Service level promised", "Date", "Required (people-h)", "Scheduled (people-h)",
            "Short (people-h)", "Covered %"] + [c for c in ("Actual calls", "Forecast error %",
                                                            "Service level achieved %") if c in cov.columns]
    df = cov[cols]
    summary = df.groupby(["Client", "Queue", "Service level promised"], as_index=False).agg(
        **{"Required (people-h)": ("Required (people-h)", "sum"), "Short (people-h)": ("Short (people-h)", "sum")})
    summary["Covered %"] = (100 * (1 - summary["Short (people-h)"] / summary["Required (people-h)"].where(
        summary["Required (people-h)"] > 0))).round(1)
    blob = workbook("Service-level report per client contract", {"Summary": summary, "By day": df},
                    context=_ctx(w), source="Live roster against the forecast requirement; actual calls and service "
                                            "level only where imported", kind="planned (and actual where imported)",
                    definitions={"Covered %": "Required people-time the roster provides; a plan figure, not measured "
                                              "service",
                                 "Service level achieved %": "Calls answered within target ÷ calls, from imported "
                                                             "actuals; blank when not imported",
                                 "Forecast error %": "Sum of half-hour absolute errors ÷ actual calls"})
    return df, blob


def print_html(title: str, sections: dict[str, pd.DataFrame], w) -> bytes:
    """A plain page to open in a browser and print: title, what it is, tables. Black on white."""
    import html
    parts = [f"<!doctype html><meta charset='utf-8'><title>{html.escape(title)}</title><style>"
             "body{font:11pt system-ui,sans-serif;margin:16mm;color:#000;background:#fff}h1{font-size:16pt}"
             "table{border-collapse:collapse;width:100%;margin:8px 0 18px}th,td{border:1px solid #999;padding:3px 6px;"
             "font-size:9pt;text-align:left}th{background:#eee}@page{size:A4 landscape;margin:12mm}"
             "p{font-size:9pt;color:#333}</style>",
             f"<h1>{html.escape(title)}</h1><p>{html.escape(_ctx(w))} · produced {datetime.now():%d %b %Y %H:%M} · "
             "all figures simulated</p>"]
    for name, df in sections.items():
        parts.append(f"<h2>{html.escape(name)}</h2>" + df.to_html(index=False, border=0, na_rep=""))
    return "".join(parts).encode()
