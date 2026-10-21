import openpyxl
import io

import pytest

from shiftsense import history, report, staffing, store
from shiftsense import workflow as wf


@pytest.fixture
def world(policy, data, req, optimized_plan, tmp_path, monkeypatch):
    monkeypatch.setenv("SHIFTSENSE_DB", str(tmp_path / "h.db"))
    return wf.new_world(policy, data, req, optimized_plan["roster"], fc=staffing.forecast(policy, data["history"]))


def test_diff_lists_before_and_after(world):
    before = world.roster.copy()
    a = world.roster[world.roster.day == world.first_day].agent_id.iloc[0]
    wf.report_absence(world, [a], world.first_day, "tl")
    d = history.diff(world.policy, before, world.roster)
    assert len(d) == 1 and d[0]["agent_id"] == a and d[0]["after"] == "Off" and d[0]["before"] != "Off"


def test_every_change_is_saved_with_details_and_snapshot(world):
    history.record_baseline(world, "Roster published", "planner")
    before, n = world.roster.copy(), len(world.audit.records)
    day = world.first_day
    inc = wf.report_absence(world, world.roster[world.roster.day == day].agent_id.tolist()[:2], day, "Rohit (TL)")
    e1 = history.record(world, before, n)
    assert e1 and e1["event"] == "AgentAbsent" and len(e1["details"]) == 2
    before, n = world.roster.copy(), len(world.audit.records)
    if inc.rec["actions"]:
        wf.approve_fix(world, inc, "Rohit (TL)", "standard cover")
        e2 = history.record(world, before, n)
        assert e2["actor"] == "Rohit (TL)" and "standard cover" in e2["reason"]
    assert history.record(world, world.roster.copy(), len(world.audit.records)) is None   # no change, no entry
    rows = history.changes(world.run_id)
    assert rows[-1]["event"] == "Roster published"
    snap = history.roster_at(rows[0]["id"])
    assert len(snap) == len(world.roster)


def test_history_survives_workspace_clear(world):
    history.record_baseline(world, "Roster published", "planner")
    store.clear()
    assert history.changes(world.run_id)


def test_weekly_report_numbers_and_excel(world):
    a = world.roster[world.roster.day == world.first_day].agent_id.iloc[0]
    wf.report_absence(world, [a], world.first_day, "tl")
    rep = report.weekly_report(world, 0)
    s = rep["summary"]
    assert 0 < s["Coverage %"] <= 100 and s["Absences"] >= 1 and s["Alerts"] >= 1
    assert set(rep["by_queue"].Queue) == set(world.policy["queues"])
    assert len(rep["attendance"]) == len(world.data["agents"])
    wb = openpyxl.load_workbook(io.BytesIO(report.to_excel(rep)))
    assert "Summary" in wb.sheetnames and "Schedule changes" in wb.sheetnames
    assert report.weekly_report(world, 1)["summary"]["Status"] == "upcoming"


def test_old_saved_workspace_is_upgraded(world):
    del world.__dict__["run_id"]
    del world.__dict__["spikes"]
    wf.upgrade(world)
    assert world.run_id and world.spikes == []
